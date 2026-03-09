#!/usr/bin/env python3
"""Generate a PPTX deck from JBehave_Basics_Presentation.md.
Usage: python docs/jBehave/generate_pptx.py
Requires: pip install -r docs/jBehave/requirements.txt

Environment variables (optional):
  MAX_BULLETS - max bullets per slide (default 8)
  TEXT_FONT_SIZE - font size for text bullets in points (default 16)
  CODE_FONT_SIZE - font size for code blocks in points (default 12)
  TEMPLATE_PATH - optional path to a reference PPTX template to use
"""
from pathlib import Path
import re
import sys
import os
import json
from datetime import datetime
from pptx import Presentation
from pptx.util import Pt

BASE_DIR = Path(__file__).resolve().parent
# Source markdown (inside docs/jBehave)
SRC = BASE_DIR / 'JBehave_Basics_Presentation.md'
# Output placed inside docs/jBehave to avoid scattering files
OUT = BASE_DIR / 'JBehave_Basics_Presentation.pptx'
CONTEXT = BASE_DIR / 'context.json'

# Template: allow user-specified via env or fall back to local template file
ENV_TEMPLATE = os.environ.get('TEMPLATE_PATH')
LOCAL_TEMPLATE = BASE_DIR / 'pptx_template.pptx'
TEMPLATE = Path(ENV_TEMPLATE) if ENV_TEMPLATE else (LOCAL_TEMPLATE if LOCAL_TEMPLATE.exists() else None)

CODE_FENCE = re.compile(r'^```')
HEADER_RE = re.compile(r'^# +')

# Config: default values (can be overridden via environment variables)
DEFAULT_MAX_BULLETS = 8
DEFAULT_TEXT_FONT_SIZE = 16
DEFAULT_CODE_FONT_SIZE = 12

# Read overrides from environment
try:
    MAX_BULLETS = int(os.environ.get('MAX_BULLETS', str(DEFAULT_MAX_BULLETS)))
except ValueError:
    MAX_BULLETS = DEFAULT_MAX_BULLETS

try:
    TEXT_FONT_SIZE_PT = int(os.environ.get('TEXT_FONT_SIZE', str(DEFAULT_TEXT_FONT_SIZE)))
except ValueError:
    TEXT_FONT_SIZE_PT = DEFAULT_TEXT_FONT_SIZE

try:
    CODE_FONT_SIZE_PT = int(os.environ.get('CODE_FONT_SIZE', str(DEFAULT_CODE_FONT_SIZE)))
except ValueError:
    CODE_FONT_SIZE_PT = DEFAULT_CODE_FONT_SIZE


def debug(msg: str):
    print(f'[generate_pptx] {msg}')


def write_context(metadata: dict):
    try:
        CONTEXT.write_text(json.dumps(metadata, indent=2), encoding='utf-8')
        debug(f'Wrote context: {CONTEXT}')
    except Exception as e:
        debug(f'Warning: failed to write context file: {e}')


def parse_markdown(md: str):
    """Parse markdown into slides with content items and speaker notes.
    Returns list of tuples: (title, items, notes_text)
    items: list of {'type': 'text'|'code', 'text': str}
    notes_text: single string (may be empty)
    """
    lines = md.splitlines()
    slides = []
    current_title = None
    current_body = []
    in_code = False
    code_buffer = []
    current_notes = []
    in_notes = False

    def append_slide(title, body, notes):
        # body: list of items
        if not body and not notes:
            return
        slides.append((title.strip(), body.copy(), notes.strip()))

    def flush():
        nonlocal current_title, current_body, current_notes, in_notes
        if current_title:
            notes_text = '\n'.join(current_notes).strip()
            # If body is long, split into multiple slides
            if current_body:
                for i in range(0, len(current_body), MAX_BULLETS):
                    chunk = current_body[i:i+MAX_BULLETS]
                    if i == 0:
                        # first slide keeps original title and notes
                        append_slide(current_title, chunk, notes_text)
                    else:
                        # continuation slides: append (cont. N)
                        cont_title = f"{current_title} (cont. {i//MAX_BULLETS + 1})"
                        # only attach notes to the first slide
                        append_slide(cont_title, chunk, "")
            else:
                append_slide(current_title, [], notes_text)
        current_title = None
        current_body = []
        current_notes = []
        in_notes = False

    for idx, line in enumerate(lines):
        if CODE_FENCE.match(line):
            # toggle code mode regardless of notes state
            if in_code:
                in_code = False
                if code_buffer:
                    code_text = '\n'.join(code_buffer)
                    # if currently in notes, treat code block as part of notes text
                    if in_notes:
                        current_notes.append(code_text)
                    else:
                        current_body.append({'type': 'code', 'text': code_text})
                    code_buffer = []
            else:
                in_code = True
            continue

        # If we are inside a code block, accumulate lines
        if in_code:
            code_buffer.append(line)
            continue

        # New slide header
        if HEADER_RE.match(line):
            if current_title:
                flush()
            current_title = line.lstrip('#').strip()
            continue

        stripped = line.strip()
        # Detect start of Notes: section (case-sensitive as in docs)
        if stripped.startswith('Notes:'):
            in_notes = True
            # capture any remaining content after 'Notes:' on same line
            tail = line.partition('Notes:')[2].strip()
            if tail:
                current_notes.append(tail)
            continue

        # If currently parsing notes, append raw line to notes
        if in_notes:
            # preserve blank lines
            current_notes.append(line)
            continue

        # Normal body text lines
        if not stripped:
            continue
        current_body.append({'type': 'text', 'text': stripped})

    if current_title:
        flush()
    return slides


def create_pptx(slides):
    # Use template if available
    if TEMPLATE:
        try:
            prs = Presentation(TEMPLATE)
            debug(f'Loaded template: {TEMPLATE}')
        except Exception as e:
            debug(f'Could not load template {TEMPLATE}: {e}; falling back to default template')
            prs = Presentation()
    else:
        prs = Presentation()

    title_layout = prs.slide_layouts[1]  # Title and Content
    code_font_size = Pt(CODE_FONT_SIZE_PT)
    text_font_size = Pt(TEXT_FONT_SIZE_PT)

    for title, items, notes in slides:
        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = title
        tf = slide.shapes.placeholders[1].text_frame
        tf.clear()
        for item in items:
            para = tf.add_paragraph()
            if item['type'] == 'code':
                para.text = 'Code:\n' + item['text']
                para.font.size = code_font_size
            else:
                txt = item['text']
                txt = txt.replace('**', '').replace('*', '')
                txt = re.sub(r'`([^`]+)`', r'\1', txt)
                para.text = txt
                para.font.size = text_font_size
            para.level = 0
        # Add speaker notes if present
        if notes:
            try:
                notes_slide = slide.notes_slide
                notes_tf = notes_slide.notes_text_frame
                # clear existing
                notes_tf.clear()
                # write notes as paragraphs
                first = True
                for line in notes.splitlines():
                    if first:
                        notes_tf.text = line
                        first = False
                    else:
                        p = notes_tf.add_paragraph()
                        p.text = line
            except Exception as e:
                debug(f'Warning: could not write notes for slide "{title}": {e}')

    prs.save(OUT)


def main():
    debug(f'Using source: {SRC}')
    if not SRC.exists():
        debug('Source markdown missing')
        raise SystemExit(f'Missing source markdown: {SRC}')
    md = SRC.read_text(encoding='utf-8')
    debug(f'Markdown length: {len(md)} chars')
    slides = parse_markdown(md)
    debug(f'Parsed {len(slides)} slides')
    if not slides:
        raise SystemExit('No slides parsed; check markdown headings.')
    create_pptx(slides)
    if OUT.exists():
        debug(f'Wrote PPTX: {OUT} ({OUT.stat().st_size} bytes)')
        # write context metadata to avoid repeating steps unnecessarily
        metadata = {
            'generated_at': datetime.utcnow().isoformat() + 'Z',
            'max_bullets': MAX_BULLETS,
            'text_font_size_pt': TEXT_FONT_SIZE_PT,
            'code_font_size_pt': CODE_FONT_SIZE_PT,
            'template_used': str(TEMPLATE) if TEMPLATE else None,
            'output_file': str(OUT)
        }
        write_context(metadata)
    else:
        debug('Failed to create PPTX file')
        raise SystemExit('PPTX not created')


if __name__ == '__main__':
    try:
        main()
    except Exception as e:
        debug(f'ERROR: {e}')
        sys.exit(1)
