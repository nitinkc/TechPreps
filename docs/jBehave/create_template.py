#!/usr/bin/env python3
"""Create a simple PPTX template file used by generate_pptx.py.
Saves `pptx_template.pptx` next to this script.
"""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pathlib import Path

OUT = Path(__file__).resolve().parent / 'pptx_template.pptx'

prs = Presentation()
# Set slide size (optional) - keep default for compatibility

# Title slide layout: adjust title font
title_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
if title:
    title.text = 'Template Title'
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(40)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)  # blue

if subtitle:
    subtitle.text = 'Template subtitle (placeholder)'
    for paragraph in subtitle.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Content slide layout: adjust placeholder font sizes/colors
content_layout = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(content_layout)
if slide2.shapes.title:
    slide2.shapes.title.text = 'Content Slide (template)'
    for paragraph in slide2.shapes.title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(28)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

# Content placeholder text
tf = slide2.shapes.placeholders[1].text_frame
tf.clear()
p = tf.add_paragraph()
p.text = '• Bullet 1 (template)'
for run in p.runs:
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
p.level = 0

p2 = tf.add_paragraph()
p2.text = '• Bullet 2 (template)'
for run in p2.runs:
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
p2.level = 0

# Save template
prs.save(OUT)
print(f'Created template: {OUT}')

